from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .aggregator import remaining_days
from .models import AnalysisResult, AssetType, UnreadableFile, ValidityRecord


RED = PatternFill("solid", fgColor="FF0000")
ORANGE = PatternFill("solid", fgColor="F4B183")
HEADER = PatternFill("solid", fgColor="173F6B")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SHEET_NAMES = {AssetType.TRACTOR: "Çekiciler", AssetType.TRAILER: "Dorseler", AssetType.ISO_TANK: "ISO Tanklar", AssetType.DRIVER: "Sürücüler"}


def _header(sheet, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[1]:
        cell.fill = HEADER
        cell.font = HEADER_FONT
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(values))}1"


def _fit(sheet) -> None:
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        sheet.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 35)


def _status(value: date, as_of: date) -> str:
    days = (value - as_of).days
    return "Geçmiş" if days < 0 else "Yaklaşıyor" if days <= 20 else "Geçerli"


def _add_asset_sheet(workbook: Workbook, title: str, rows, as_of: date) -> None:
    sheet = workbook.create_sheet(title)
    documents = sorted({name for row in rows for name in row.document_dates})
    _header(sheet, ["Nakliyeci", "Araç / Kimlik", "En Son Geliş Tarihi", "Geliş Sayısı", *documents])
    for row in rows:
        values = [row.carrier, row.asset_id, row.latest_control_date, row.arrival_count]
        values.extend(row.document_dates.get(name) for name in documents)
        sheet.append(values)
        for index, name in enumerate(documents, start=5):
            cell = sheet.cell(sheet.max_row, index)
            if cell.value:
                status = _status(cell.value, as_of)
                if status == "Geçmiş": cell.fill = RED
                elif status == "Yaklaşıyor": cell.fill = ORANGE
    _fit(sheet)


def write_report(result: AnalysisResult, unreadable: list[UnreadableFile], output_path: Path, as_of: date) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Özet"
    _header(summary, ["Araç Türü", "Benzersiz Kayıt", "Kritik Evrak"])
    for asset_type in AssetType:
        critical = sum(1 for item in result.critical if item.asset_type is asset_type)
        summary.append([asset_type.value, len(result.by_type.get(asset_type, [])), critical])
    _fit(summary)
    for asset_type, title in SHEET_NAMES.items():
        _add_asset_sheet(workbook, title, result.by_type.get(asset_type, []), as_of)
    critical_sheet = workbook.create_sheet("Kritik Evraklar")
    _header(critical_sheet, ["Nakliyeci", "Araç Türü", "Araç / Kimlik", "Soru / Evrak", "Geçerlilik Tarihi", "Son Kontrol Tarihi", "Kalan Gün", "Durum"])
    for record in result.critical:
        days = remaining_days(record, as_of)
        status = "Geçmiş" if days < 0 else "Yaklaşıyor"
        critical_sheet.append([record.carrier, record.asset_type.value, record.asset_id, record.short_name, record.validity_date, record.control_date, days, status])
        for cell in critical_sheet[critical_sheet.max_row]:
            cell.fill = RED if status == "Geçmiş" else ORANGE
    _fit(critical_sheet)
    errors = workbook.create_sheet("Okunamayan Dosyalar")
    _header(errors, ["Dosya", "Hata Nedeni"])
    for item in unreadable:
        errors.append([item.source_path, item.reason])
    _fit(errors)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
