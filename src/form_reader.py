from __future__ import annotations

from datetime import date, datetime
import re

from openpyxl import load_workbook

from .models import ArchiveFile, AssetType, UnreadableFile, ValidityRecord
from .question_catalog import QuestionCatalog


HEADERS = {"sorumlu", "soru", "geçerlilik tarihi", "tarih"}


def _header_key(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _datetime(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    for fmt in (
        "%d.%m.%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S",
        "%d.%m.%Y %H:%M", "%d/%m/%Y %H:%M",
        "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(str(value).strip(), fmt)
        except ValueError:
            pass
    return None


def _date(value: object) -> date | None:
    parsed = _datetime(value)
    return parsed.date() if parsed else None


def _asset_type(identifier: str, question: str) -> AssetType:
    lower_question = question.lower()
    compact = re.sub(r"\s+", "", identifier).upper()
    if re.fullmatch(r"\d{11}", compact):
        return AssetType.DRIVER
    if "dorse" in lower_question or "römork" in lower_question:
        return AssetType.TRAILER
    if "iso" in lower_question or "tank konteyner" in lower_question or re.fullmatch(r"[A-Z]{4}\d{7}", compact):
        return AssetType.ISO_TANK
    return AssetType.TRACTOR


def _columns(sheet) -> tuple[int, dict[str, int]] | None:
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        headers = {_header_key(value): index for index, value in enumerate(row)}
        if HEADERS <= headers.keys():
            return row_number, headers
    return None


def read_form(file: ArchiveFile, catalog: QuestionCatalog) -> tuple[list[ValidityRecord], list[UnreadableFile]]:
    try:
        sheet = load_workbook(file.path, data_only=True, read_only=True).active
        found = _columns(sheet)
        if not found:
            return [], [UnreadableFile(str(file.path), "Sorumlu/Soru/Geçerlilik Tarihi/Tarih başlıkları bulunamadı")]
        header_row, columns = found
        records: list[ValidityRecord] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            question = str(row[columns["soru"]] or "").strip()
            identifier = str(row[columns["sorumlu"]] or "").strip()
            validity = _date(row[columns["geçerlilik tarihi"]])
            if not question or not identifier or validity is None:
                continue
            control_at = _datetime(row[columns["tarih"]]) or datetime.combine(file.archive_date, datetime.min.time())
            records.append(ValidityRecord(
                carrier=file.carrier, asset_type=_asset_type(identifier, question), asset_id=identifier,
                question=question, short_name=catalog.short_name(question), validity_date=validity,
                control_date=control_at.date(), control_at=control_at, source_path=str(file.path),
            ))
        return records, []
    except Exception as error:
        return [], [UnreadableFile(str(file.path), str(error))]
