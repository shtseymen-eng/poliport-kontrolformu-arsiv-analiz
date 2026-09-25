import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from src.aggregator import summarize
from src.excel_reporter import write_report
from src.models import AssetType, ValidityRecord


class ExcelReporterTests(unittest.TestCase):
    def test_report_has_asset_sheets_and_colored_critical_rows(self):
        record = ValidityRecord("AGENA LOJ.", AssetType.TRACTOR, "34ABC123", "Muayene", "Muayene", date(2026, 9, 20), date(2026, 9, 1), "f.xlsx")
        with TemporaryDirectory() as temp:
            output = Path(temp) / "rapor.xlsx"
            write_report(summarize([record], date(2026, 9, 25)), [], output, date(2026, 9, 25))
            book = load_workbook(output)
            self.assertTrue({"Özet", "Çekiciler", "Dorseler", "ISO Tanklar", "Sürücüler", "Kritik Evraklar", "Okunamayan Dosyalar"} <= set(book.sheetnames))
            self.assertEqual(book["Kritik Evraklar"][2][7].value, "Geçmiş")
